#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.14"
# dependencies = [
#   "cyclopts",
#   "openpyxl",
#   "requests",
# ]
# ///
"""Post HPC usage reminders from the team's planning sheet to Slack."""

import io
import re
import time
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import openpyxl
import requests

NZ = ZoneInfo("Pacific/Auckland")

# Each tab of the planning sheet and the columns the bot reads from it.
TABS = {
    "Runs": ["Person", "Description", "HPC", "Start date", "End date"],
    "People": ["Name", "Slack member ID", "Admin"],
    "HPCs": ["HPC", "Alert if idle"],
}
MEMBER_ID = re.compile(r"[UW][A-Z0-9]{6,}")


# ---------------------------------------------------------------- dates


def week_monday(day: date) -> date:
    """Return the Monday of the Monday-to-Sunday week that contains ``day``."""
    return day - timedelta(days=day.weekday())


def _span(start: date, end: date, first: str = "", last: str = "") -> str:
    """Format two different dates, naming the month and year only as often as needed."""
    if start.year != end.year:
        return f"{first}{start.day} {start:%b %Y} – {last}{end.day} {end:%b %Y}"
    if start.month != end.month:
        return f"{first}{start.day} {start:%b} – {last}{end.day} {end:%b}"
    return f"{first}{start.day} – {last}{end.day} {end:%b}"


def format_date_range(start: date, end: date) -> str:
    """Format an inclusive date range, e.g. ``12 Oct – 1 Nov``."""
    if start == end:
        return f"{start.day} {start:%b}"
    return _span(start, end)


def format_week(monday: date) -> str:
    """Format the week starting ``monday``, e.g. ``Mon 19 – Sun 25 Oct``."""
    return _span(monday, monday + timedelta(days=6), "Mon ", "Sun ")


def nz_today(now: datetime | None = None) -> date:
    """Return today's date in New Zealand (``now`` must be timezone-aware)."""
    return (now or datetime.now(NZ)).astimezone(NZ).date()


# ---------------------------------------------------------------- the sheet


class SheetError(Exception):
    """The planning sheet couldn't be downloaded or opened."""


@dataclass(frozen=True)
class Person:
    name: str
    member_id: str
    admin: bool


@dataclass(frozen=True)
class Hpc:
    name: str
    watched: bool  # "Alert if idle" is ticked


@dataclass(frozen=True)
class Run:
    row: int  # the sheet's own row number (header = 1)
    person: Person
    description: str
    hpc: str
    start: date
    end: date  # inclusive


@dataclass(frozen=True)
class Schedule:
    """The usable contents of the sheet, plus a line for everything that wasn't usable."""

    people: list[Person]
    hpcs: list[Hpc]
    runs: list[Run]
    problems: list[str]

    @property
    def admins(self) -> list[Person]:
        return [person for person in self.people if person.admin]


def read_workbook(xlsx: bytes) -> dict[str, list[tuple]]:
    """Return every tab of an .xlsx file as a list of row tuples, header row included."""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
        try:
            return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}
        finally:
            workbook.close()
    except Exception as exc:
        raise SheetError(f"couldn't open the downloaded spreadsheet: {exc}") from exc


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _key(value: object) -> str:
    """Normalise a name for matching, ignoring case and surrounding spaces."""
    return _text(value).casefold()


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _records(
    tabs: Mapping[str, Sequence[Sequence[object]]], problems: list[str]
) -> dict[str, list[tuple[int, dict[str, object]]]]:
    """Return each tab's data rows as (sheet row number, {column: value}).

    A missing tab or column is added to ``problems`` and that tab is treated as empty.
    """
    by_key = {_key(title): rows for title, rows in tabs.items()}
    records = {}
    for tab, columns in TABS.items():
        rows = by_key.get(_key(tab))
        if rows is None:
            problems.append(f'The sheet has no "{tab}" tab.')
            records[tab] = []
            continue
        header = [_key(cell) for cell in rows[0]] if rows else []
        missing = [column for column in columns if _key(column) not in header]
        if missing:
            problems.extend(f'The {tab} tab has no "{column}" column.' for column in missing)
            records[tab] = []
            continue
        index = {column: header.index(_key(column)) for column in columns}
        records[tab] = [
            (number, {column: row[i] if i < len(row) else None for column, i in index.items()})
            for number, row in enumerate(rows[1:], start=2)
        ]
    return records


def parse_schedule(tabs: Mapping[str, Sequence[Sequence[object]]], this_monday: date) -> Schedule:
    """Turn the sheet's tabs into a Schedule, skipping and reporting anything unusable.

    Parameters
    ----------
    tabs
        Tab title to rows, as returned by `read_workbook`.
    this_monday
        Monday of the current week. Runs that ended before it are ignored unchecked.
    """
    problems: list[str] = []
    records = _records(tabs, problems)

    people: dict[str, Person] = {}
    seen: set[str] = set()
    for number, record in records["People"]:
        name = _text(record["Name"])
        if not name:
            continue
        if _key(name) in seen:
            problems.append(f'People row {number}: "{name}" is already listed above, so this row was skipped.')
            continue
        seen.add(_key(name))
        member_id = _text(record["Slack member ID"])
        if not member_id:
            problems.append(f"People row {number} ({name}): no Slack member ID, so this row was skipped.")
        elif not MEMBER_ID.fullmatch(member_id):
            problems.append(f'People row {number} ({name}): "{member_id}" isn\'t a Slack member ID, so this row was skipped.')
        else:
            people[_key(name)] = Person(name, member_id, record["Admin"] is True)
    if not any(person.admin for person in people.values()):
        problems.append("No admin is ticked on the People tab.")

    hpcs: dict[str, Hpc] = {}
    for number, record in records["HPCs"]:
        name = _text(record["HPC"])
        if not name:
            continue
        if _key(name) in hpcs:
            problems.append(f'HPCs row {number}: "{name}" is already listed above, so this row was skipped.')
            continue
        hpcs[_key(name)] = Hpc(name, record["Alert if idle"] is True)

    runs: list[Run] = []
    for number, record in records["Runs"]:
        if not any(_text(record[column]) for column in TABS["Runs"]):
            continue
        start, end = _as_date(record["Start date"]), _as_date(record["End date"])
        if end is not None and end < this_monday:
            continue
        person_text, hpc_text = _text(record["Person"]), _text(record["HPC"])
        person, hpc = people.get(_key(person_text)), hpcs.get(_key(hpc_text))
        reasons = []
        if not person_text:
            reasons.append("no person given")
        elif person is None:
            reasons.append(f'"{person_text}" isn\'t on the People tab with a valid Slack member ID')
        if not hpc_text:
            reasons.append("no HPC given")
        elif hpc is None:
            reasons.append(f'HPC "{hpc_text}" isn\'t on the HPCs tab')
        for label, column, parsed in (("start date", "Start date", start), ("end date", "End date", end)):
            if not _text(record[column]):
                reasons.append(f"{label} is empty")
            elif parsed is None:
                reasons.append(f'{label} "{_text(record[column])}" isn\'t a date')
        if start and end and start > end:
            reasons.append("start date is after the end date")
        if reasons:
            problems.append(f"Row {number} ({person_text or '?'}, {hpc_text or '?'}): {'; '.join(reasons)}, so it was skipped.")
        else:
            runs.append(Run(number, person, _text(record["Description"]), hpc.name, start, end))

    return Schedule(list(people.values()), list(hpcs.values()), runs, problems)


# ---------------------------------------------------------------- choosing runs


def overlaps(run: Run, monday: date) -> bool:
    """Return True if the run is on during any day of the week starting ``monday``."""
    return run.start <= monday + timedelta(days=6) and run.end >= monday


def _by_start(runs) -> list[Run]:
    return sorted(runs, key=lambda run: (run.start, run.person.name.casefold(), run.row))


def runs_in_week(schedule: Schedule, monday: date) -> list[Run]:
    """Runs that overlap the week starting ``monday``, earliest start first."""
    return _by_start(run for run in schedule.runs if overlaps(run, monday))


def runs_starting_in_week(schedule: Schedule, monday: date) -> list[Run]:
    """Runs whose start date falls in the week starting ``monday``, earliest first."""
    return _by_start(run for run in schedule.runs if monday <= run.start <= monday + timedelta(days=6))


def idle_hpcs(schedule: Schedule, monday: date) -> list[str]:
    """Watched HPCs that no run overlaps in the week starting ``monday``, in HPCs-tab order."""
    covered = {run.hpc for run in runs_in_week(schedule, monday)}
    return [hpc.name for hpc in schedule.hpcs if hpc.watched and hpc.name not in covered]


# ---------------------------------------------------------------- message text


def sheet_url(sheet_id: str) -> str:
    """Link to the planning sheet, shown at the end of each post."""
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}/edit"


def escape(text: str) -> str:
    """Escape the characters Slack's mrkdwn uses for mentions and links."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _mentions(people: list[Person]) -> str:
    return " ".join(f"<@{person.member_id}>" for person in people)


def _run_groups(schedule: Schedule, runs: list[Run]) -> list[str]:
    """List ``runs`` under a bold heading for each HPC, in HPCs-tab order."""
    lines = []
    for hpc in schedule.hpcs:
        group = [run for run in runs if run.hpc == hpc.name]
        if group:
            lines.append(f"*{escape(hpc.name)}*")
            for run in group:
                about = f" — {escape(run.description)}" if run.description else ""
                lines.append(f"• <@{run.person.member_id}>{about} ({format_date_range(run.start, run.end)})")
    return lines


def _sheet_link_line(sheet_link: str) -> str:
    return f"Plans changed? <{sheet_link}|Update the sheet>"


def format_weekly_post(schedule: Schedule, monday: date, sheet_link: str) -> str:
    """The Monday post: this week's runs by HPC, then an idle-HPC alert for this week and next."""
    next_monday = monday + timedelta(days=7)
    sections = [
        f":spiral_calendar_pad: *HPC plan for this week* · {format_week(monday)}",
        "\n".join(_run_groups(schedule, runs_in_week(schedule, monday))) or "Nothing planned this week.",
    ]
    idle_now, idle_next = idle_hpcs(schedule, monday), idle_hpcs(schedule, next_monday)
    if idle_now or idle_next:
        mentions = _mentions(schedule.admins)
        alert = [f":warning: {mentions} nothing is planned on:" if mentions else ":warning: Nothing is planned on:"]
        alert += [f"• *{escape(name)}* this week" for name in idle_now]
        next_week = format_date_range(next_monday, next_monday + timedelta(days=6))
        alert += [f"• *{escape(name)}* next week ({next_week})" for name in idle_next]
        sections.append("\n".join(alert))
    sections.append(_sheet_link_line(sheet_link))
    return "\n\n".join(sections)


def format_heads_up_post(schedule: Schedule, next_monday: date, sheet_link: str) -> str | None:
    """The Friday post: runs that start next week, or None if nothing starts."""
    groups = _run_groups(schedule, runs_starting_in_week(schedule, next_monday))
    if not groups:
        return None
    return "\n\n".join([
        f":eyes: *Heads-up: starting next week* · {format_week(next_monday)}",
        "\n".join(groups),
        _sheet_link_line(sheet_link),
    ])


def format_problem_reply(schedule: Schedule) -> str | None:
    """The thread reply listing problems in the sheet, or None if there are none."""
    count = len(schedule.problems)
    if not count:
        return None
    header = "1 problem in the sheet needs fixing:" if count == 1 else f"{count} problems in the sheet need fixing:"
    if mentions := _mentions(schedule.admins):
        header = f"{mentions} {header}"
    return "\n".join([header, *(f"• {escape(problem)}" for problem in schedule.problems)])


def format_sheet_error(what: str, error: str, sheet_link: str) -> str:
    """The channel post sent instead of ``what`` when the sheet can't be read."""
    return f":x: Couldn't read the HPC planning sheet, so {what} didn't go out: {escape(error)}\n<{sheet_link}|Open the sheet>"


# ---------------------------------------------------------------- HTTP


class SlackError(Exception):
    """A message couldn't be posted to Slack."""


class _Transient(Exception):
    """An HTTP response worth retrying (429 or 5xx)."""

    def __init__(self, status_code: int):
        super().__init__(f"HTTP {status_code}")
        self.status_code = status_code


def _checked(response: requests.Response) -> requests.Response:
    if response.status_code == 429 or response.status_code >= 500:
        raise _Transient(response.status_code)
    return response


def with_retries[T](call: Callable[[], T], attempts: int = 3, wait_seconds: float = 5.0) -> T:
    """Run ``call``, retrying network failures and 429/5xx responses."""
    for attempt in range(1, attempts + 1):
        try:
            return call()
        except (requests.ConnectionError, requests.Timeout, _Transient):
            if attempt == attempts:
                raise
            time.sleep(wait_seconds)
    raise AssertionError("attempts must be at least 1")


def download_workbook(sheet_id: str) -> bytes:
    """Download the planning sheet as .xlsx, without credentials (it is link-shared)."""
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    try:
        response = with_retries(lambda: _checked(requests.get(url, timeout=30)))
    except _Transient as exc:
        raise SheetError(f"download failed with HTTP {exc.status_code}") from exc
    except requests.RequestException as exc:
        raise SheetError(f"download failed: {exc}") from exc
    if response.status_code != 200:
        raise SheetError(f"download failed with HTTP {response.status_code}")
    if "spreadsheetml" not in response.headers.get("Content-Type", ""):
        raise SheetError("the download wasn't a spreadsheet; check the sheet is shared as \"Anyone with the link can view\"")
    return response.content


def post_message(token: str, channel: str, text: str, thread_ts: str | None = None) -> str:
    """Post ``text`` to a Slack channel (or reply in a thread) and return the message's ts."""
    payload = {"channel": channel, "text": text, "unfurl_links": False, "unfurl_media": False}
    if thread_ts is not None:
        payload["thread_ts"] = thread_ts
    headers = {"Authorization": f"Bearer {token}"}
    try:
        response = with_retries(lambda: _checked(requests.post(
            "https://slack.com/api/chat.postMessage", headers=headers, json=payload, timeout=30)))
    except _Transient as exc:
        raise SlackError(f"Slack returned HTTP {exc.status_code}") from exc
    except requests.RequestException as exc:
        raise SlackError(f"couldn't reach Slack: {exc}") from exc
    data = response.json()
    if not data.get("ok"):
        raise SlackError(f"Slack said: {data.get('error')}")
    return data["ts"]
